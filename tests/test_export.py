"""Automated test suite for CSV, Excel, and PDF timetable exporting."""

import csv
import tempfile
from pathlib import Path
import pytest
import openpyxl

from core.models import ScheduledSession, Timetable
from exchange.exporters import (
    CSVExporter,
    ExcelExporter,
    PDFTimetableExporter,
)
from exchange.models import (
    ConflictReportRow,
    ExportFormat,
    RoomUtilizationRow,
    TimetableLayout,
    UnscheduledSessionRow,
    WorkloadReportRow,
)
from exchange.service import ImportExportServiceImpl


@pytest.fixture
def temp_dir():
    with tempfile.TemporaryDirectory() as td:
        yield Path(td)


@pytest.fixture
def sample_sessions():
    """Generates synthetic scheduled sessions for testing."""
    return [
        ScheduledSession(
            id="sess-1",
            timetable_id="tt-01",
            class_subject_assignment_id="csa-1",
            day_of_week=0,  # Monday
            start_period=0,  # Period 1
            duration_periods=1,
            room_id="room-101",
            teacher_id="t-alan",
            class_group_id="class-3a",
        ),
        ScheduledSession(
            id="sess-2",
            timetable_id="tt-01",
            class_subject_assignment_id="csa-2",
            day_of_week=0,  # Monday
            start_period=1,  # Period 2
            duration_periods=2,  # 2-period lab block
            room_id="room-lab",
            teacher_id="t-ada",
            class_group_id="class-3a",
        ),
        ScheduledSession(
            id="sess-3",
            timetable_id="tt-01",
            class_subject_assignment_id="csa-3",
            day_of_week=1,  # Tuesday
            start_period=2,  # Period 3
            duration_periods=1,
            room_id="room-101",
            teacher_id="t-alan",
            class_group_id="class-3b",
        ),
    ]


@pytest.fixture
def sample_metadata():
    return {
        "institution_name": "Turing Polytechnic Institute",
        "department_name": "Computer Science & Engineering",
        "academic_year_name": "2026-2027",
        "semester_name": "Fall 2026",
        "timetable_name": "Odd Semester Master Schedule",
        "version": 1,
        "is_published": False,
        "days": [0, 1, 2, 3, 4],
        "periods_per_day": 6,
        "period_times": [("09:00", "10:00"), ("10:00", "11:00"), ("11:15", "12:15"), ("13:00", "14:00"), ("14:00", "15:00"), ("15:00", "16:00")],
        "class_name_map": {"class-3a": "CS-3A", "class-3b": "CS-3B"},
        "sub_code_map": {"class-3a": "CS301", "class-3b": "CS302"},
        "sub_name_map": {"class-3a": "Data Structures, Algorithms & Advanced Complexity", "class-3b": "Database Management Systems"},
        "teacher_emp_map": {"t-alan": "T101", "t-ada": "T102"},
        "teacher_name_map": {"t-alan": "Dr. Alan Turing", "t-ada": "Prof. Ada Lovelace"},
        "room_code_map": {"room-101": "R-101", "room-lab": "LAB-AI"},
        "room_name_map": {"room-101": "Hall 101", "room-lab": "AI Computing Facility"},
    }


# ---------------------------------------------------------------------------
# CSV Export Tests
# ---------------------------------------------------------------------------


def test_export_sessions_csv(temp_dir, sample_sessions, sample_metadata):
    csv_out = temp_dir / "sessions.csv"
    res = CSVExporter.export_sessions(sample_sessions, csv_out, sample_metadata)

    assert res.success is True
    assert res.format == ExportFormat.CSV
    assert res.records_exported == 3
    assert csv_out.exists()

    with open(csv_out, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        rows = list(reader)

    assert len(rows) == 3
    assert rows[0]["session_id"] == "sess-1"
    assert rows[0]["day_name"] == "Monday"
    assert rows[0]["class_name"] == "CS-3A"
    assert rows[0]["teacher_name"] == "Dr. Alan Turing"
    assert rows[1]["duration_periods"] == "2"


def test_export_empty_sessions_csv(temp_dir):
    csv_out = temp_dir / "empty.csv"
    res = CSVExporter.export_sessions([], csv_out)
    assert res.success is True
    assert res.records_exported == 0
    assert csv_out.exists()


def test_export_csv_escapes_commas_and_quotes(temp_dir):
    sessions = [
        ScheduledSession(
            id="sess-special",
            timetable_id="tt-01",
            class_subject_assignment_id="csa-1",
            day_of_week=0,
            start_period=0,
            duration_periods=1,
            room_id="room-1",
            teacher_id="t-1",
            class_group_id="class-1",
        )
    ]
    meta = {
        "class_name_map": {"class-1": 'Cohort "A", Special'},
        "sub_name_map": {"class-1": "Intro to CS, Math & Logic"},
    }
    csv_out = temp_dir / "escaped.csv"
    res = CSVExporter.export_sessions(sessions, csv_out, meta)
    assert res.success is True

    # Re-read and assert exact string preserved
    with open(csv_out, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        row = next(reader)
    assert row["class_name"] == 'Cohort "A", Special'
    assert row["subject_name"] == "Intro to CS, Math & Logic"


# ---------------------------------------------------------------------------
# Excel Export Tests
# ---------------------------------------------------------------------------


def test_export_excel_workbook(temp_dir, sample_sessions, sample_metadata):
    xlsx_out = temp_dir / "timetable_full.xlsx"
    w_rows = [
        WorkloadReportRow(
            teacher_id="t-alan",
            teacher_name="Dr. Alan Turing",
            employee_id="T101",
            department_name="CSE",
            assigned_periods=14,
            assigned_contact_hours=14.0,
            min_periods=10,
            max_periods=18,
            remaining_capacity=4,
            max_consecutive_observed=2,
            is_overloaded=False,
            is_underloaded=False,
            has_workload_rule=True,
            warning_message="Normal",
        )
    ]
    r_rows = [
        RoomUtilizationRow(
            room_id="room-101",
            room_code="R-101",
            room_name="Hall 101",
            room_type="LECTURE_HALL",
            capacity=60,
            building="Main",
            scheduled_sessions_count=2,
            occupied_periods=2,
            available_periods=30,
            utilization_rate=6.67,
            notes="Normal",
        )
    ]

    res = ExcelExporter.export_full_workbook(
        output_path=xlsx_out,
        sessions=sample_sessions,
        workload_rows=w_rows,
        room_rows=r_rows,
        metadata=sample_metadata,
    )

    assert res.success is True
    assert res.format == ExportFormat.EXCEL
    assert xlsx_out.exists()

    # Programmatically inspect generated workbook structure
    wb = openpyxl.load_workbook(xlsx_out)
    sheet_names = wb.sheetnames
    assert "Master Schedule" in sheet_names
    assert "Class Schedules" in sheet_names
    assert "Faculty Schedules" in sheet_names
    assert "Room Occupancy" in sheet_names
    assert "Faculty Workload" in sheet_names
    assert "Room Utilization" in sheet_names

    # Check Master Schedule sheet contents
    ws_master = wb["Master Schedule"]
    assert ws_master.cell(row=1, column=1).value == "Day"
    assert ws_master.cell(row=2, column=1).value == "Monday"
    assert ws_master.freeze_panes == "A2"

    wb.close()


def test_export_empty_excel(temp_dir):
    xlsx_out = temp_dir / "empty.xlsx"
    res = ExcelExporter.export_full_workbook(output_path=xlsx_out)
    assert res.success is True
    assert xlsx_out.exists()


# ---------------------------------------------------------------------------
# PDF Export Tests
# ---------------------------------------------------------------------------


def test_export_pdf_timetable(temp_dir, sample_sessions, sample_metadata):
    pdf_out = temp_dir / "timetable.pdf"
    res = PDFTimetableExporter.export_timetable(
        output_path=pdf_out,
        sessions=sample_sessions,
        layout=TimetableLayout.CLASS_GRID,
        metadata=sample_metadata,
    )
    assert res.success is True
    assert res.format == ExportFormat.PDF
    assert pdf_out.exists()
    assert pdf_out.stat().st_size > 1000

    # Verify PDF magic bytes
    with open(pdf_out, "rb") as f:
        magic = f.read(5)
    assert magic == b"%PDF-"


def test_export_pdf_teacher_and_room_layouts(temp_dir, sample_sessions, sample_metadata):
    t_pdf = temp_dir / "teacher_timetable.pdf"
    res_t = PDFTimetableExporter.export_timetable(
        output_path=t_pdf,
        sessions=sample_sessions,
        layout=TimetableLayout.TEACHER_GRID,
        metadata=sample_metadata,
    )
    assert res_t.success is True
    assert t_pdf.exists()

    r_pdf = temp_dir / "room_timetable.pdf"
    res_r = PDFTimetableExporter.export_timetable(
        output_path=r_pdf,
        sessions=sample_sessions,
        layout=TimetableLayout.ROOM_GRID,
        metadata=sample_metadata,
    )
    assert res_r.success is True
    assert r_pdf.exists()


def test_export_pdf_empty_dataset(temp_dir, sample_metadata):
    empty_pdf = temp_dir / "empty.pdf"
    res = PDFTimetableExporter.export_timetable(
        output_path=empty_pdf,
        sessions=[],
        metadata=sample_metadata,
    )
    assert res.success is True
    assert empty_pdf.exists()


# ---------------------------------------------------------------------------
# Service-Level Integration Export Tests
# ---------------------------------------------------------------------------


def test_service_export_methods(temp_dir, sample_sessions, sample_metadata):
    service = ImportExportServiceImpl(context={"sessions": sample_sessions, "metadata": sample_metadata})

    csv_path = str(temp_dir / "service_export.csv")
    res_csv = service.export_timetable_to_csv("tt-01", csv_path)
    assert Path(res_csv).exists()

    xlsx_path = str(temp_dir / "service_export.xlsx")
    res_xlsx = service.export_timetable_to_excel("tt-01", xlsx_path)
    assert Path(res_xlsx).exists()

    pdf_path = str(temp_dir / "service_export.pdf")
    res_pdf = service.export_timetable_to_pdf("tt-01", pdf_path)
    assert Path(res_pdf).exists()
