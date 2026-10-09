"""Generates downloadable CSV and Excel (.xlsx) templates for all academic imports.

Includes:
- Standard column headers aligned with import validators.
- Built-in documentation and instructions.
- Clear required vs optional indications.
- Valid fictional sample records.
"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Dict, List, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from exchange.models import DatasetType
from exchange.security import validate_export_file_path

# Styling constants
HEADER_FILL_REQ = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FILL_OPT = PatternFill(start_color="5B9BD5", end_color="5B9BD5", fill_type="solid")
HEADER_FONT = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
INSTR_TITLE_FONT = Font(name="Segoe UI", size=14, bold=True, color="1F4E79")
INSTR_BODY_FONT = Font(name="Segoe UI", size=10, color="333333")
INSTR_BOLD_FONT = Font(name="Segoe UI", size=10, bold=True, color="1F4E79")
SAMPLE_FONT = Font(name="Segoe UI", size=10, italic=False, color="000000")

THIN_SIDE = Side(border_style="thin", color="D9D9D9")
BOX_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)

# Template catalog definitions
# (DatasetType, description, columns [(name, is_required, description)], sample_rows)
TEMPLATE_CATALOG: Dict[DatasetType, Dict[str, Any]] = {
    DatasetType.DEPARTMENTS: {
        "title": "Academic Departments Import Template",
        "description": "Defines institutional departments. Import this first before subjects, faculty, and classes.",
        "order": 1,
        "columns": [
            ("code", True, "Unique department code (e.g., CSE, MECH, EEE, MATH)"),
            ("name", True, "Full official department name (e.g., Computer Science & Engineering)"),
            ("institution_code", False, "Institution identifier (optional, default: DEFAULT)"),
        ],
        "samples": [
            ["CSE", "Computer Science & Engineering", "MAIN"],
            ["ECE", "Electronics & Communication Engineering", "MAIN"],
            ["MATH", "Mathematics & Basic Sciences", "MAIN"],
        ],
    },
    DatasetType.ACADEMIC_PERIODS: {
        "title": "Academic Years & Semesters Import Template",
        "description": "Defines active academic year cycles and numbered semesters.",
        "order": 2,
        "columns": [
            ("academic_year", True, "Academic cycle identifier (e.g., 2026-2027)"),
            ("start_date", False, "Cycle start date in YYYY-MM-DD format"),
            ("end_date", False, "Cycle end date in YYYY-MM-DD format"),
            ("is_current", False, "Active cycle flag (TRUE/FALSE)"),
            ("semester_number", True, "Numeric semester index (e.g., 1, 2, 3, 4, 5, 6, 7, 8)"),
            ("term_name", False, "Human-readable semester term (e.g., Fall 2026, Odd Semester)"),
        ],
        "samples": [
            ["2026-2027", "2026-08-01", "2027-05-31", "TRUE", 1, "Semester 1 (Fall 2026)"],
            ["2026-2027", "2026-08-01", "2027-05-31", "TRUE", 3, "Semester 3 (Fall 2026)"],
            ["2026-2027", "2026-08-01", "2027-05-31", "TRUE", 5, "Semester 5 (Fall 2026)"],
        ],
    },
    DatasetType.CLASSES: {
        "title": "Class Cohorts & Sections Import Template",
        "description": "Defines student class sections. Requires Department code.",
        "order": 3,
        "columns": [
            ("name", True, "Unique class cohort name (e.g., CS-3A, S3AD)"),
            ("department_code", True, "Referenced department code (must already exist)"),
            ("student_count", False, "Total enrolled students in cohort (default: 60)"),
            ("semester_number", False, "Semester number (default: 1)"),
            ("section", False, "Section letter (e.g., A, B)"),
            ("home_room_code", False, "Default home classroom code if assigned"),
        ],
        "samples": [
            ["CS-3A", "CSE", 60, 3, "A", "R-101"],
            ["CS-3B", "CSE", 55, 3, "B", "R-102"],
            ["EC-3A", "ECE", 50, 3, "A", "R-201"],
        ],
    },
    DatasetType.ROOMS: {
        "title": "Physical Rooms & Facilities Import Template",
        "description": "Defines classrooms, laboratories, and lecture halls.",
        "order": 4,
        "columns": [
            ("code", True, "Unique room identifier or number (e.g., R-101, LAB-AI)"),
            ("name", True, "Display name (e.g., Lecture Hall 101, Advanced AI Computing Lab)"),
            ("room_type", False, "Type: LECTURE_HALL, COMPUTER_LAB, HARDWARE_LAB, SEMINAR_ROOM, WORKSHOP"),
            ("capacity", True, "Maximum seated student capacity (integer >= 1)"),
            ("building", False, "Campus building or block name"),
            ("floor", False, "Floor number (integer)"),
        ],
        "samples": [
            ["R-101", "Lecture Hall 101", "LECTURE_HALL", 65, "Main Academic Block", 1],
            ["R-102", "Classroom 102", "LECTURE_HALL", 60, "Main Academic Block", 1],
            ["LAB-AI", "AI & Data Computing Lab", "COMPUTER_LAB", 40, "Turing Technology Centre", 2],
        ],
    },
    DatasetType.TEACHERS: {
        "title": "Faculty & Teachers Import Template",
        "description": "Defines teacher profiles and contractual teaching hours.",
        "order": 5,
        "columns": [
            ("employee_id", True, "Unique institutional faculty ID (e.g., T101, EMP-042)"),
            ("name", True, "Full faculty name with title (e.g., Dr. Alan Turing)"),
            ("department_code", True, "Referenced department code (must already exist)"),
            ("email", False, "Institutional email address"),
            ("is_full_time", False, "Full-time status (TRUE/FALSE, default: TRUE)"),
            ("max_weekly_hours", False, "Max permitted weekly teaching periods (default: 20)"),
            ("min_weekly_hours", False, "Min required weekly teaching periods (default: 0)"),
            ("max_daily_hours", False, "Max permitted teaching periods per day (default: 5)"),
            ("max_consecutive_hours", False, "Max consecutive periods without break (default: 3)"),
        ],
        "samples": [
            ["T101", "Dr. Alan Turing", "CSE", "alan.turing@institution.edu", "TRUE", 18, 10, 4, 2],
            ["T102", "Prof. Ada Lovelace", "CSE", "ada.lovelace@institution.edu", "TRUE", 16, 8, 4, 2],
            ["T103", "Dr. Claude Shannon", "ECE", "claude.shannon@institution.edu", "TRUE", 20, 12, 5, 3],
        ],
    },
    DatasetType.SUBJECTS: {
        "title": "Master Subject Catalog Import Template",
        "description": "Defines course catalog with credits and required contact hours.",
        "order": 6,
        "columns": [
            ("code", True, "Unique course code (e.g., CS301, EC204)"),
            ("name", True, "Official subject name (e.g., Data Structures & Algorithms)"),
            ("department_code", True, "Owning department code"),
            ("credits", False, "Academic credits (degree weight, default: 3.0)"),
            ("weekly_contact_hours", False, "Physical classroom hours needed per week (default: 3.0)"),
            ("is_elective", False, "Elective flag (TRUE/FALSE, default: FALSE)"),
            ("requires_lab", False, "Requires laboratory facility (TRUE/FALSE, default: FALSE)"),
        ],
        "samples": [
            ["CS301", "Data Structures & Algorithms", "CSE", 4.0, 4.0, "FALSE", "FALSE"],
            ["CS331", "Data Structures Laboratory", "CSE", 2.0, 3.0, "FALSE", "TRUE"],
            ["MA301", "Discrete Mathematics", "MATH", 4.0, 4.0, "FALSE", "FALSE"],
        ],
    },
    DatasetType.TEACHER_QUALIFICATIONS: {
        "title": "Faculty Subject Eligibility Import Template",
        "description": "Specifies which subjects each teacher is qualified and prefers to teach.",
        "order": 7,
        "columns": [
            ("employee_id", True, "Teacher employee ID"),
            ("subject_code", True, "Subject course code"),
            ("competency_level", False, "1=Primary competency, 2=Secondary/Backup (default: 1)"),
            ("preference_score", False, "Faculty preference score 1 to 5 (5=Highest, default: 5)"),
        ],
        "samples": [
            ["T101", "CS301", 1, 5],
            ["T101", "CS331", 1, 4],
            ["T102", "CS301", 2, 3],
        ],
    },
    DatasetType.SUBJECT_OFFERINGS: {
        "title": "Semester Subject Offerings Import Template",
        "description": "Specifies weekly session requirements for subjects offered in a semester.",
        "order": 8,
        "columns": [
            ("subject_code", True, "Subject course code"),
            ("semester_number", False, "Semester number (default: 1)"),
            ("weekly_contact_hours", True, "Total weekly contact hours required"),
            ("weekly_periods", True, "Total discrete weekly timetable periods required"),
        ],
        "samples": [
            ["CS301", 3, 4.0, 4],
            ["CS331", 3, 3.0, 3],
            ["MA301", 3, 4.0, 4],
        ],
    },
    DatasetType.CLASS_ASSIGNMENTS: {
        "title": "Class-Subject Faculty Assignments Import Template",
        "description": "Links a class cohort, subject, and designated primary teacher.",
        "order": 9,
        "columns": [
            ("class_name", True, "Class cohort name (e.g., CS-3A)"),
            ("subject_code", True, "Subject course code (e.g., CS301)"),
            ("primary_teacher_employee_id", True, "Assigned primary faculty employee ID"),
            ("secondary_teacher_employee_id", False, "Optional co-teacher for labs/practicals"),
            ("student_group_name", False, "Optional specific lab batch (e.g., Batch 1)"),
        ],
        "samples": [
            ["CS-3A", "CS301", "T101", "", ""],
            ["CS-3A", "CS331", "T101", "T102", "Batch 1"],
            ["CS-3A", "MA301", "T103", "", ""],
        ],
    },
    DatasetType.TEACHER_AVAILABILITY: {
        "title": "Teacher Availability & Preferences Import Template",
        "description": "Declares unworkable timeslots (hard constraint) and slot preferences (-10 to +10).",
        "order": 10,
        "columns": [
            ("employee_id", True, "Teacher employee ID"),
            ("day_of_week", True, "Day index: 0=Mon, 1=Tue, 2=Wed, 3=Thu, 4=Fri (or name 'Monday')"),
            ("period_index", True, "Period index (0-indexed: 0..periods-1)"),
            ("is_available", False, "TRUE=Available, FALSE=Blocked slot (default: TRUE)"),
            ("preference_weight", False, "Weight -10 (strongly avoid) to +10 (strongly prefer, default: 0)"),
        ],
        "samples": [
            ["T101", 4, 5, "FALSE", -10],  # Friday Period 5 unavailable
            ["T101", 0, 0, "TRUE", 5],     # Monday Period 0 strongly preferred
            ["T102", 2, 0, "FALSE", -5],    # Wednesday morning blocked
        ],
    },
}


class TemplateGenerator:
    """Generates standard CSV and Excel templates."""

    @classmethod
    def generate_csv_template(cls, dataset_type: DatasetType, output_path: str | Path) -> Path:
        """Writes a clean CSV template with headers and fictional sample rows."""
        dest = validate_export_file_path(output_path, expected_ext=".csv")
        info = TEMPLATE_CATALOG.get(dataset_type)
        if not info:
            raise ValueError(f"Unknown dataset type: {dataset_type}")

        fieldnames = [col[0] for col in info["columns"]]
        samples = info["samples"]

        with open(dest, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(fieldnames)
            for row in samples:
                writer.writerow(row)

        return dest

    @classmethod
    def generate_excel_template(cls, dataset_type: DatasetType, output_path: str | Path) -> Path:
        """Writes a rich Excel template with Instructions and Data worksheets."""
        dest = validate_export_file_path(output_path, expected_ext=".xlsx")
        info = TEMPLATE_CATALOG.get(dataset_type)
        if not info:
            raise ValueError(f"Unknown dataset type: {dataset_type}")

        wb = openpyxl.Workbook()

        # 1. Instructions Sheet
        ws_instr = wb.active
        ws_instr.title = "Instructions"
        ws_instr.views.sheetView[0].showGridLines = True

        ws_instr.cell(row=1, column=1, value=info["title"]).font = INSTR_TITLE_FONT
        ws_instr.cell(row=2, column=1, value=info["description"]).font = INSTR_BODY_FONT
        ws_instr.cell(row=3, column=1, value=f"Recommended import sequence order: Step {info['order']}").font = INSTR_BOLD_FONT

        ws_instr.cell(row=5, column=1, value="Column Specifications:").font = INSTR_BOLD_FONT
        headers_instr = ["Column Name", "Requirement", "Data Type & Guidance"]
        for c_idx, h in enumerate(headers_instr, start=1):
            cell = ws_instr.cell(row=6, column=c_idx, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL_REQ
            cell.border = BOX_BORDER

        for r_idx, (col_name, is_req, col_desc) in enumerate(info["columns"], start=7):
            req_str = "REQUIRED (*)" if is_req else "Optional"
            c1 = ws_instr.cell(row=r_idx, column=1, value=col_name)
            c2 = ws_instr.cell(row=r_idx, column=2, value=req_str)
            c3 = ws_instr.cell(row=r_idx, column=3, value=col_desc)
            for c in (c1, c2, c3):
                c.font = INSTR_BODY_FONT
                c.border = BOX_BORDER
            if is_req:
                c2.font = Font(name="Segoe UI", size=10, bold=True, color="C00000")

        # Instructions note
        curr_row = 7 + len(info["columns"]) + 2
        ws_instr.cell(row=curr_row, column=1, value="Instructions for use:").font = INSTR_BOLD_FONT
        notes = [
            "1. Switch to the 'Data' worksheet to enter your institutional records.",
            "2. Columns marked with (*) are mandatory and cannot be left blank.",
            "3. Dark Blue headers indicate REQUIRED fields; Light Blue headers indicate OPTIONAL fields.",
            "4. Sample rows are provided as valid illustrative examples; replace them with your actual data.",
            "5. Save the completed file as .xlsx or export as .csv before importing.",
        ]
        for note in notes:
            curr_row += 1
            ws_instr.cell(row=curr_row, column=1, value=note).font = INSTR_BODY_FONT

        # Autofit columns in instructions sheet
        for col in ws_instr.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws_instr.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 4, 18)

        # 2. Data Sheet
        ws_data = wb.create_sheet(title="Data")
        ws_data.views.sheetView[0].showGridLines = True

        for c_idx, (col_name, is_req, _) in enumerate(info["columns"], start=1):
            display_name = f"{col_name} *" if is_req else col_name
            cell = ws_data.cell(row=1, column=c_idx, value=display_name)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL_REQ if is_req else HEADER_FILL_OPT
            cell.border = BOX_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")

        ws_data.row_dimensions[1].height = 24

        # Add sample data rows
        for r_idx, row_values in enumerate(info["samples"], start=2):
            for c_idx, val in enumerate(row_values, start=1):
                cell = ws_data.cell(row=r_idx, column=c_idx, value=val)
                cell.font = SAMPLE_FONT
                cell.border = BOX_BORDER

        # Autofit data sheet
        for col in ws_data.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws_data.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 5, 16)

        ws_data.freeze_panes = "A2"

        wb.save(dest)
        wb.close()
        return dest

    @classmethod
    def generate_all_templates(cls, directory_path: str | Path) -> List[Path]:
        """Generates both CSV and Excel templates for every dataset into target directory."""
        dir_p = Path(directory_path)
        dir_p.mkdir(parents=True, exist_ok=True)
        created_files: List[Path] = []

        for ds_type in DatasetType:
            slug = ds_type.value.lower()
            csv_path = dir_p / f"{slug}_template.csv"
            xlsx_path = dir_p / f"{slug}_template.xlsx"
            cls.generate_csv_template(ds_type, csv_path)
            cls.generate_excel_template(ds_type, xlsx_path)
            created_files.extend([csv_path, xlsx_path])

        return created_files
