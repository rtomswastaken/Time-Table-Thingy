"""Rich, formatted Excel (.xlsx) workbook generator using openpyxl.

Produces multi-sheet workbooks with professional typography, auto-adjusted
column widths, frozen headers, styled weekly grid views, and formula protection.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.models import ScheduledSession
from exchange.models import (
    ConflictReportRow,
    ExportFormat,
    ExportResult,
    RoomUtilizationRow,
    UnscheduledSessionRow,
    WorkloadReportRow,
)
from exchange.security import sanitize_spreadsheet_value, validate_export_file_path

DAY_NAMES = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]

# Styling palettes
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Segoe UI", size=14, bold=True, color="1F4E79")
SUBTITLE_FONT = Font(name="Segoe UI", size=10, italic=True, color="595959")
GRID_HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
GRID_HEADER_FONT = Font(name="Segoe UI", size=10, bold=True, color="000000")
CELL_FONT = Font(name="Segoe UI", size=10, color="000000")
BOLD_FONT = Font(name="Segoe UI", size=10, bold=True, color="000000")
WARN_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
ERROR_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
SESSION_CELL_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

THIN_BORDER_SIDE = Side(border_style="thin", color="D9D9D9")
BOX_BORDER = Border(
    left=THIN_BORDER_SIDE,
    right=THIN_BORDER_SIDE,
    top=THIN_BORDER_SIDE,
    bottom=THIN_BORDER_SIDE,
)


def autofit_columns(sheet: openpyxl.worksheet.worksheet.Worksheet, max_col_width: int = 50) -> None:
    """Calculates maximum string length per column and applies generous column width."""
    for col in sheet.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if "\n" in val:
                val = max(val.split("\n"), key=len)
            max_len = max(max_len, len(val))
        sheet.column_dimensions[col_letter].width = min(max_len + 4, max_col_width)


class ExcelExporter:
    """Exports timetable schedules and administrative reports to styled .xlsx workbooks."""

    @classmethod
    def export_full_workbook(
        cls,
        output_path: str | Path,
        sessions: Optional[List[ScheduledSession]] = None,
        workload_rows: Optional[List[WorkloadReportRow]] = None,
        room_rows: Optional[List[RoomUtilizationRow]] = None,
        conflict_rows: Optional[List[ConflictReportRow]] = None,
        unscheduled_rows: Optional[List[UnscheduledSessionRow]] = None,
        metadata: Optional[Dict[str, Any]] = None,
    ) -> ExportResult:
        """Constructs an integrated, multi-tab institutional timetable spreadsheet."""
        dest = validate_export_file_path(output_path, expected_ext=".xlsx")
        wb = openpyxl.Workbook()
        # Remove default empty sheet
        default_sheet = wb.active
        meta = metadata or {}

        sheets_created = 0

        # 1. Master Sessions Sheet
        if sessions is not None and len(sessions) > 0:
            cls._build_master_sessions_sheet(wb, sessions, meta)
            sheets_created += 1

        # 2. Class Timetable Grids Sheet
        if sessions is not None and len(sessions) > 0:
            cls._build_grid_sheet(wb, "Class Schedules", sessions, group_by="class", meta=meta)
            sheets_created += 1

        # 3. Faculty Timetable Grids Sheet
        if sessions is not None and len(sessions) > 0:
            cls._build_grid_sheet(wb, "Faculty Schedules", sessions, group_by="teacher", meta=meta)
            sheets_created += 1

        # 4. Room Occupancy Grids Sheet
        if sessions is not None and len(sessions) > 0:
            cls._build_grid_sheet(wb, "Room Occupancy", sessions, group_by="room", meta=meta)
            sheets_created += 1

        # 5. Workload Sheet
        if workload_rows is not None and len(workload_rows) > 0:
            cls._build_workload_sheet(wb, workload_rows)
            sheets_created += 1

        # 6. Room Utilization Sheet
        if room_rows is not None and len(room_rows) > 0:
            cls._build_room_utilization_sheet(wb, room_rows)
            sheets_created += 1

        # 7. Conflicts Sheet
        if conflict_rows is not None and len(conflict_rows) > 0:
            cls._build_conflicts_sheet(wb, conflict_rows)
            sheets_created += 1

        # 8. Unscheduled Sessions Sheet
        if unscheduled_rows is not None and len(unscheduled_rows) > 0:
            cls._build_unscheduled_sheet(wb, unscheduled_rows)
            sheets_created += 1

        if sheets_created == 0:
            # Create a placeholder empty schedule sheet
            ws = wb.create_sheet(title="Timetable")
            ws.cell(row=1, column=1, value="No timetable data or reports were provided for export.")
            ws.cell(row=1, column=1).font = TITLE_FONT

        # Remove initial default sheet if unused
        if default_sheet and default_sheet.title in wb.sheetnames and len(wb.sheetnames) > 1:
            wb.remove(default_sheet)

        wb.save(dest)
        wb.close()

        return ExportResult(
            success=True,
            format=ExportFormat.EXCEL,
            file_path=str(dest),
            file_size_bytes=dest.stat().st_size,
            records_exported=len(sessions or []),
            message=f"Exported {sheets_created} sheets to {dest.name}.",
        )

    @classmethod
    def _build_master_sessions_sheet(
        cls,
        wb: openpyxl.Workbook,
        sessions: List[ScheduledSession],
        meta: Dict[str, Any],
    ) -> None:
        ws = wb.create_sheet(title="Master Schedule")
        ws.views.sheetView[0].showGridLines = True

        headers = [
            "Day",
            "Period",
            "Class",
            "Subject Code",
            "Subject Name",
            "Faculty ID",
            "Faculty Name",
            "Room Code",
            "Room Name",
            "Duration",
            "Batch / Group",
            "Locked",
        ]

        # Header row
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BOX_BORDER
        ws.row_dimensions[1].height = 24

        class_name_map = meta.get("class_name_map", {})
        sub_code_map = meta.get("sub_code_map", {})
        sub_name_map = meta.get("sub_name_map", {})
        teacher_emp_map = meta.get("teacher_emp_map", {})
        teacher_name_map = meta.get("teacher_name_map", {})
        room_code_map = meta.get("room_code_map", {})
        room_name_map = meta.get("room_name_map", {})

        # Sort sessions chronologically: Day, Period, Class
        sorted_sessions = sorted(sessions, key=lambda s: (s.day_of_week, s.start_period, s.class_group_id))

        for row_idx, s in enumerate(sorted_sessions, start=2):
            day_name = DAY_NAMES[s.day_of_week] if 0 <= s.day_of_week < len(DAY_NAMES) else f"Day {s.day_of_week}"
            row_data = [
                day_name,
                f"Period {s.start_period + 1}",
                sanitize_spreadsheet_value(class_name_map.get(s.class_group_id, s.class_group_id)),
                sanitize_spreadsheet_value(sub_code_map.get(s.class_group_id, "")),
                sanitize_spreadsheet_value(sub_name_map.get(s.class_group_id, "")),
                sanitize_spreadsheet_value(teacher_emp_map.get(s.teacher_id, s.teacher_id)),
                sanitize_spreadsheet_value(teacher_name_map.get(s.teacher_id, "")),
                sanitize_spreadsheet_value(room_code_map.get(s.room_id, s.room_id)),
                sanitize_spreadsheet_value(room_name_map.get(s.room_id, "")),
                s.duration_periods,
                sanitize_spreadsheet_value(s.student_group_id or "Whole Class"),
                "Yes" if s.is_locked else "No",
            ]
            for col_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = CELL_FONT
                cell.border = BOX_BORDER
                if col_idx in (1, 2, 10, 12):
                    cell.alignment = Alignment(horizontal="center")

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(sorted_sessions) + 1}"
        autofit_columns(ws)

    @classmethod
    def _build_grid_sheet(
        cls,
        wb: openpyxl.Workbook,
        sheet_title: str,
        sessions: List[ScheduledSession],
        group_by: str,
        meta: Dict[str, Any],
    ) -> None:
        """Constructs formatted weekly grid matrices stacked by entity (Class, Teacher, or Room)."""
        ws = wb.create_sheet(title=sheet_title)
        ws.views.sheetView[0].showGridLines = True

        days = meta.get("days", [0, 1, 2, 3, 4])
        periods_per_day = meta.get("periods_per_day", 6)

        class_name_map = meta.get("class_name_map", {})
        sub_code_map = meta.get("sub_code_map", {})
        teacher_emp_map = meta.get("teacher_emp_map", {})
        teacher_name_map = meta.get("teacher_name_map", {})
        room_code_map = meta.get("room_code_map", {})

        # Group sessions by entity key
        grouped: Dict[str, List[ScheduledSession]] = {}
        for s in sessions:
            if group_by == "class":
                key = s.class_group_id
            elif group_by == "teacher":
                key = s.teacher_id
            else:
                key = s.room_id
            grouped.setdefault(key, []).append(s)

        current_row = 1

        for entity_id, entity_sessions in sorted(grouped.items()):
            # Entity Header
            if group_by == "class":
                display_title = f"Class: {class_name_map.get(entity_id, entity_id)}"
            elif group_by == "teacher":
                t_name = teacher_name_map.get(entity_id, "")
                t_code = teacher_emp_map.get(entity_id, entity_id)
                display_title = f"Faculty: {t_name} ({t_code})" if t_name else f"Faculty: {t_code}"
            else:
                display_title = f"Room: {room_code_map.get(entity_id, entity_id)}"

            title_cell = ws.cell(row=current_row, column=1, value=display_title)
            title_cell.font = TITLE_FONT
            current_row += 1

            # Grid header: Day | Period 1 | Period 2 | ...
            ws.cell(row=current_row, column=1, value="Day / Period").fill = GRID_HEADER_FILL
            ws.cell(row=current_row, column=1).font = GRID_HEADER_FONT
            ws.cell(row=current_row, column=1).border = BOX_BORDER
            ws.cell(row=current_row, column=1).alignment = Alignment(horizontal="center", vertical="center")

            for p in range(periods_per_day):
                cell = ws.cell(row=current_row, column=p + 2, value=f"Period {p + 1}")
                cell.fill = GRID_HEADER_FILL
                cell.font = GRID_HEADER_FONT
                cell.border = BOX_BORDER
                cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[current_row].height = 20
            current_row += 1

            # Map placements: (day, period) -> string
            matrix: Dict[tuple[int, int], List[str]] = {}
            for s in entity_sessions:
                for period_offset in range(s.duration_periods):
                    slot = (s.day_of_week, s.start_period + period_offset)
                    # Label
                    parts = []
                    sub = sub_code_map.get(s.class_group_id, "")
                    if sub:
                        parts.append(sub)
                    if group_by != "class":
                        c_name = class_name_map.get(s.class_group_id, "")
                        if c_name:
                            parts.append(c_name)
                    if group_by != "teacher":
                        t_code = teacher_emp_map.get(s.teacher_id, "")
                        if t_code:
                            parts.append(t_code)
                    if group_by != "room":
                        r_code = room_code_map.get(s.room_id, "")
                        if r_code:
                            parts.append(r_code)

                    desc = " / ".join(parts) if parts else "Session"
                    matrix.setdefault(slot, []).append(desc)

            # Output days rows
            for d in days:
                d_name = DAY_NAMES[d] if 0 <= d < len(DAY_NAMES) else f"Day {d}"
                day_cell = ws.cell(row=current_row, column=1, value=d_name)
                day_cell.font = BOLD_FONT
                day_cell.border = BOX_BORDER
                day_cell.alignment = Alignment(horizontal="center", vertical="center")

                for p in range(periods_per_day):
                    slot_content = matrix.get((d, p), [])
                    text_val = "\n".join(slot_content)
                    cell = ws.cell(row=current_row, column=p + 2, value=text_val)
                    cell.font = CELL_FONT
                    cell.border = BOX_BORDER
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    if slot_content:
                        cell.fill = SESSION_CELL_FILL

                ws.row_dimensions[current_row].height = 36
                current_row += 1

            # Blank spacing row between entities
            current_row += 2

        autofit_columns(ws, max_col_width=25)

    @classmethod
    def _build_workload_sheet(cls, wb: openpyxl.Workbook, rows: List[WorkloadReportRow]) -> None:
        ws = wb.create_sheet(title="Faculty Workload")
        ws.views.sheetView[0].showGridLines = True

        headers = [
            "Faculty ID",
            "Faculty Name",
            "Department",
            "Weekly Periods",
            "Contact Hours",
            "Min Required",
            "Max Permitted",
            "Remaining Capacity",
            "Max Consecutive",
            "Overloaded",
            "Underloaded",
            "Notes / Warnings",
        ]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.border = BOX_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 24

        for row_idx, r in enumerate(rows, start=2):
            cells_data = [
                sanitize_spreadsheet_value(r.employee_id),
                sanitize_spreadsheet_value(r.teacher_name),
                sanitize_spreadsheet_value(r.department_name),
                r.assigned_periods,
                r.assigned_contact_hours,
                r.min_periods if r.min_periods is not None else "N/A",
                r.max_periods if r.max_periods is not None else "N/A",
                r.remaining_capacity if r.remaining_capacity is not None else "N/A",
                r.max_consecutive_observed,
                "YES" if r.is_overloaded else "No",
                "YES" if r.is_underloaded else "No",
                sanitize_spreadsheet_value(r.warning_message),
            ]
            for col_idx, val in enumerate(cells_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = CELL_FONT
                cell.border = BOX_BORDER
                if r.is_overloaded:
                    cell.fill = ERROR_FILL
                elif r.is_underloaded:
                    cell.fill = WARN_FILL

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        autofit_columns(ws)

    @classmethod
    def _build_room_utilization_sheet(cls, wb: openpyxl.Workbook, rows: List[RoomUtilizationRow]) -> None:
        ws = wb.create_sheet(title="Room Utilization")
        ws.views.sheetView[0].showGridLines = True

        headers = [
            "Room Code",
            "Room Name",
            "Type",
            "Capacity",
            "Building",
            "Sessions Scheduled",
            "Occupied Periods",
            "Available Periods",
            "Utilization Rate (%)",
            "Notes",
        ]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.border = BOX_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 24

        for row_idx, r in enumerate(rows, start=2):
            cells_data = [
                sanitize_spreadsheet_value(r.room_code),
                sanitize_spreadsheet_value(r.room_name),
                sanitize_spreadsheet_value(r.room_type),
                r.capacity,
                sanitize_spreadsheet_value(r.building),
                r.scheduled_sessions_count,
                r.occupied_periods,
                r.available_periods,
                round(r.utilization_rate, 1),
                sanitize_spreadsheet_value(r.notes),
            ]
            for col_idx, val in enumerate(cells_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = CELL_FONT
                cell.border = BOX_BORDER

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
        autofit_columns(ws)

    @classmethod
    def _build_conflicts_sheet(cls, wb: openpyxl.Workbook, rows: List[ConflictReportRow]) -> None:
        ws = wb.create_sheet(title="Scheduling Conflicts")
        ws.views.sheetView[0].showGridLines = True

        headers = [
            "Constraint Type",
            "Severity",
            "Description",
            "Involved Faculty",
            "Involved Rooms",
            "Involved Classes",
            "Remediation Guidance",
        ]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.border = BOX_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 24

        for row_idx, r in enumerate(rows, start=2):
            cells_data = [
                sanitize_spreadsheet_value(r.violation_type),
                r.severity,
                sanitize_spreadsheet_value(r.description),
                "; ".join(r.involved_teachers),
                "; ".join(r.involved_rooms),
                "; ".join(r.involved_classes),
                sanitize_spreadsheet_value(r.remediation_suggestion),
            ]
            for col_idx, val in enumerate(cells_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = CELL_FONT
                cell.border = BOX_BORDER
                if r.is_hard:
                    cell.fill = ERROR_FILL
                else:
                    cell.fill = WARN_FILL

        ws.freeze_panes = "A2"
        autofit_columns(ws, max_col_width=40)

    @classmethod
    def _build_unscheduled_sheet(cls, wb: openpyxl.Workbook, rows: List[UnscheduledSessionRow]) -> None:
        ws = wb.create_sheet(title="Unscheduled Sessions")
        ws.views.sheetView[0].showGridLines = True

        headers = [
            "Subject Code",
            "Subject Name",
            "Class Cohort",
            "Assigned Faculty",
            "Required Meetings",
            "Placed Meetings",
            "Unplaced Meetings",
            "Diagnostic Reason",
        ]
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.border = BOX_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 24

        for row_idx, r in enumerate(rows, start=2):
            cells_data = [
                sanitize_spreadsheet_value(r.subject_code),
                sanitize_spreadsheet_value(r.subject_name),
                sanitize_spreadsheet_value(r.class_name),
                sanitize_spreadsheet_value(r.teacher_name),
                r.required_sessions,
                r.placed_sessions,
                r.missing_count,
                sanitize_spreadsheet_value(r.reason),
            ]
            for col_idx, val in enumerate(cells_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = CELL_FONT
                cell.border = BOX_BORDER
                cell.fill = WARN_FILL

        ws.freeze_panes = "A2"
        autofit_columns(ws)
